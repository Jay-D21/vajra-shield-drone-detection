import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Set inner padding for table cells in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def format_table(table, col_widths=None, header_bg="E2E8F0", alt_bg="F8FAFC"):
    """Format table with clean styling, borders, and alternating row shading."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    
    for r_idx, row in enumerate(table.rows):
        is_header = (r_idx == 0)
        for c_idx, cell in enumerate(row.cells):
            set_cell_margins(cell, top=130, bottom=130, left=170, right=170)
            if is_header:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.bold = True
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(15, 23, 42)
            else:
                bg = alt_bg if (r_idx % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(30, 41, 59)
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]

def add_heading_with_spacing(doc, text, level, space_before=12, space_after=6):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(space_before)
    h.paragraph_format.space_after = Pt(space_after)
    h.paragraph_format.keep_with_next = True
    return h

def add_figure_with_caption(doc, img_path, caption_text, width=Inches(5.8)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(14)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(71, 85, 105)
    else:
        print(f"Warning: Image path not found: {img_path}")

print("Helper functions defined successfully.")
